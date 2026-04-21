"""Script export service — generate TXT, Markdown, JSON, Word formats."""

import io
import json
import re
from typing import Any, Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from loguru import logger


def _html_to_plain(html: str) -> str:
    """Strip HTML tags, decode common entities, and normalise whitespace."""
    text = re.sub(r"<br\s*/?>", "\n", html, flags=re.IGNORECASE)
    text = re.sub(r"</p>|</h[1-6]>|</li>|<hr\s*/?>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    text = text.replace("&nbsp;", " ").replace("&quot;", '"').replace("&#39;", "'")
    # Collapse more-than-two consecutive blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


class ScriptExportService:
    """Export script chapters to TXT, Markdown, JSON, or DOCX."""

    # ─── Public API ──────────────────────────────────────────────────────

    def export_txt(self, project: Dict[str, Any], chapters: List[Dict]) -> str:
        """Return plain-text representation of the script."""
        lines = [project.get("name", "Untitled")]
        genre = project.get("genre")
        if genre:
            lines.append(f"Type: {genre}")
        lines.append("")

        for ch in chapters:
            num = ch.get("chapter_number", "?")
            title = ch.get("title", "")
            lines.append(f"Chapter {num}  {title}")
            lines.append("")

            content = self._chapter_body(ch)
            if content:
                lines.append(content)
                lines.append("")

            summary = ch.get("summary", "")
            if summary:
                lines.append(f"Summary: {summary}")
                lines.append("")

        result = "\n".join(lines)
        logger.debug("TXT export: %d chars", len(result))
        return result

    def export_markdown(self, project: Dict[str, Any], chapters: List[Dict]) -> str:
        """Return Markdown representation of the script."""
        lines = [f"# {project.get('name', 'Untitled')}"]
        genre = project.get("genre")
        if genre:
            lines.append(f"**Type**: {genre}")
        lines.append("")
        lines.append("## Content")
        lines.append("")

        for ch in chapters:
            num = ch.get("chapter_number", "?")
            title = ch.get("title", "")
            lines.append(f"### Chapter {num}  {title}")
            lines.append("")

            content = self._chapter_body(ch)
            if content:
                lines.append(content)
                lines.append("")

            summary = ch.get("summary", "")
            if summary:
                lines.append(f"> Summary: {summary}")
                lines.append("")

        result = "\n".join(lines)
        logger.debug("Markdown export: %d chars", len(result))
        return result

    def export_json(
        self, project: Dict[str, Any], chapters: List[Dict]
    ) -> Dict[str, Any]:
        """Return a serialisable dict representing the full script."""
        return {
            "name": project.get("name", "Untitled"),
            "genre": project.get("genre"),
            "description": project.get("description"),
            "chapters": [
                {
                    "number": ch.get("chapter_number"),
                    "title": ch.get("title"),
                    "summary": ch.get("summary"),
                    "content_text": self._chapter_body(ch),
                }
                for ch in chapters
            ],
        }

    def export_docx(self, project: Dict[str, Any], chapters: List[Dict]) -> io.BytesIO:
        """Build and return a DOCX document as a BytesIO buffer."""
        doc = Document()

        # Title
        title_para = doc.add_heading(project.get("name", "Untitled"), level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Genre
        genre = project.get("genre")
        if genre:
            p = doc.add_paragraph(f"Type: {genre}")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph("")

        for ch in chapters:
            num = ch.get("chapter_number", "?")
            title = ch.get("title", "")
            doc.add_heading(f"Chapter {num}  {title}", level=2)

            content = self._chapter_body(ch)
            if content:
                for para_text in content.split("\n"):
                    stripped = para_text.strip()
                    if stripped:
                        doc.add_paragraph(stripped)

            summary = ch.get("summary", "")
            if summary:
                p = doc.add_paragraph(f"Summary: {summary}")
                for run in p.runs:
                    run.italic = True

            doc.add_paragraph("")

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        logger.debug("DOCX export: %d chapters", len(chapters))
        return buffer

    # ─── Helpers ─────────────────────────────────────────────────────────

    def _chapter_body(self, chapter: Dict) -> str:
        """Return the best available plain-text content for a chapter.

        Priority: content_html (HTML → plain) > content (plain) > "".
        """
        html = chapter.get("content_html", "")
        if html and html.strip():
            return _html_to_plain(html)
        return (chapter.get("content") or "").strip()
